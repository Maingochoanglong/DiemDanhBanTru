"""Điều phối các thao tác của app (không phụ thuộc giao diện)."""

from __future__ import annotations

import datetime as dt
import json
import os
from dataclasses import dataclass, field
from pathlib import Path

from .core import matcher
from .core.day import DayState
from .core.meals import MealLayout, MealSummary
from .core.models import ClassInfo
from .core.parser import ParsedMessage, parse_messages
from .core.text import class_sort_key, fold
from .excel import sobook
from .excel.database import AppData, default_data_path, short_codes
from .excel.files import is_locked
from .excel.roster_import import apply_changes, diff_roster, read_cdhs

APP_NAME = "DiemDanhBanTru"


def config_path() -> Path:
    base = os.environ.get("APPDATA") or str(Path.home() / ".config")
    return Path(base) / APP_NAME / "config.json"


def load_config() -> dict:
    try:
        return json.loads(config_path().read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def save_config(cfg: dict) -> None:
    p = config_path()
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")


def next_school_day(day: dt.date, step: int = 1) -> dt.date:
    day += dt.timedelta(days=step)
    while day.weekday() >= 5:
        day += dt.timedelta(days=step)
    return day


@dataclass
class PasteOutcome:
    message: ParsedMessage
    results: list[matcher.MatchResult] = field(default_factory=list)
    guessed_from: list[str] = field(default_factory=list)  # các lớp đoán được khi tin không ghi lớp
    class_by: str = ""  # "sender": lớp theo tên người gửi; "names": theo tên học sinh
    class_how: str = ""  # mô tả cách đoán lớp cho giáo vụ, vd. "theo tên Zalo Bích Nguyễn"
    suggestions: list = field(default_factory=list)  # đề nghị thêm/cắt bán trú (tin báo lịch)
    schedule_known: bool = False  # tin báo lịch nhưng lịch đó đã có sẵn


class NeedMonthBook(Exception):
    """Chưa có sổ cho tháng của ngày đang chọn."""

    def __init__(self, month: int, year: int, prev: Path | None):
        super().__init__(f"Chưa có sổ tháng {month}/{year}")
        self.month, self.year, self.prev = month, year, prev


class Workspace:
    def __init__(self, data_path: str | Path | None = None):
        cfg = load_config()
        self.data = AppData.load(data_path or cfg.get("data_path") or default_data_path())
        self.day: dt.date = dt.date.today()
        self.book_path: Path | None = None
        self.state: DayState | None = None
        self.layout: MealLayout = MealLayout()
        self.sheet_checks: dict[str, list[matcher.MatchResult]] = {}
        self.template_name: str | None = None
        self._undo: list[DayState] = []
        self._wb_cache: tuple[tuple, object] | None = None
        self.sheet_name: str = ""  # tên sheet của ngày đang mở (vd. "01" hoặc "29")
        self.warning: str = ""  # cảnh báo sau lần lưu gần nhất
        self.dirty = False  # ngày đang mở có thay đổi chưa ghi vào sổ
        self.data_dirty = False  # danh sách HS / ghép lớp có thay đổi chưa lưu
        self.paste_skipped: dict[dt.date, int] = {}  # tin Zalo của ngày khác (lần dán gần nhất)
        self.unplaced: list[ParsedMessage] = []  # tin có tên vắng nhưng không tìm được lớp (chưa xử lý)

    # ---------------------------------------------------------------- dữ liệu
    @property
    def roster(self):
        return self.data.roster

    @property
    def settings(self) -> dict[str, str]:
        return self.data.settings

    @property
    def needs_setup(self) -> bool:
        return not self.roster.students

    def remember_data_path(self) -> None:
        cfg = load_config()
        cfg["data_path"] = str(self.data.path)
        save_config(cfg)

    def save_data(self) -> None:
        self.roster.invalidate()
        self.data.save()
        self.remember_data_path()
        self.data_dirty = False

    def import_cdhs(self, path: str | Path):
        """Trả về (các thay đổi, các dòng đọc được) để giao diện cho duyệt."""
        rows = read_cdhs(path)
        return diff_roster(self.roster, rows), rows

    def apply_import(self, changes, cdhs_path: str | Path) -> int:
        n = apply_changes(self.roster, changes)
        self.settings["cdhs_file"] = str(cdhs_path)
        self._ensure_classes()
        self.save_data()
        return n

    def _ensure_classes(self, gvcn: dict[str, str] | None = None) -> None:
        codes = sorted({s.lop_tinh for s in self.roster.students if s.ban_tru} | set(self.roster.classes), key=class_sort_key)
        for code, ma in short_codes(codes).items():
            info = self.roster.classes.setdefault(code, ClassInfo(code))
            if not info.ma_ngan:
                info.ma_ngan = ma
        for code, name in (gvcn or {}).items():
            info = self.roster.classes.setdefault(code, ClassInfo(code))
            if name and not info.gvcn:
                info.gvcn = name

    # ---------------------------------------------------------------- sổ tháng
    def set_book(self, path: str | Path) -> None:
        path = Path(path)
        self.settings["so_file"] = str(path)
        self.settings["so_folder"] = str(path.parent)
        self.save_data()

    def _book(self, path: str | Path):
        """Đọc sổ (dùng lại bản đã đọc nếu file chưa thay đổi). Chỉ để đọc, không sửa bản này."""
        key = (str(path), Path(path).stat().st_mtime_ns)
        if self._wb_cache and self._wb_cache[0] == key:
            return self._wb_cache[1]
        wb = sobook.load_book(path)
        self._wb_cache = (key, wb)
        return wb

    def book_for(self, day: dt.date) -> Path | None:
        cur = self.settings.get("so_file")
        if cur and Path(cur).exists():
            wb = self._book(cur)
            bm = sobook.book_month(wb, cur)
            if bm is None or bm == (day.month, day.year):
                return Path(cur)
        folder = self.settings.get("so_folder")
        if folder:
            found = sobook.find_month_file(folder, day.month, day.year)
            if found:
                self.settings["so_file"] = str(found)
                return found
        return None

    def create_month_book(self, day: dt.date) -> Path:
        prev = self.settings.get("so_file")
        if not prev or not Path(prev).exists():
            raise FileNotFoundError("Chưa chọn sổ tháng trước để làm mẫu.")
        new = sobook.month_file_name(Path(prev), day.month)
        if new.exists():
            raise FileExistsError(str(new))
        sobook.new_month_book(prev, day, new)
        self.set_book(new)
        return new

    def page1_area(self, info: sobook.SheetInfo | None) -> str:
        if self.settings.get("vung_in_1"):
            return self.settings["vung_in_1"]
        return "A1:G46" if info is None else f"A1:{_col(info.cols['ten'])}{info.total_row + 6}"

    def page2_area(self) -> str:
        return self.settings.get("vung_in_2") or self.layout.area

    # ---------------------------------------------------------------- ngày
    def load_day(self, day: dt.date) -> DayState:
        """Mở ngày: nếu sổ đã có sheet ngày đó thì đọc lại và kiểm tra; không thì lấy sheet trước làm mẫu."""
        self.day = day
        self._undo.clear()
        self.dirty = False
        self.sheet_checks = {}
        self.unplaced = []
        path = self.book_for(day)
        self.book_path = path
        if path is None:
            prev = Path(self.settings["so_file"]) if self.settings.get("so_file") else None
            self.state = DayState.new(day, self.roster, self.roster.class_codes())
            self.layout = MealLayout()
            raise NeedMonthBook(day.month, day.year, prev)

        wb = self._book(path)
        sheets = sobook.day_sheets(wb)
        ws = sheets.get(day.day)
        if ws is not None and sobook.sheet_date(ws) not in (None, day):
            ws = None  # sheet trùng số ngày nhưng khác tháng
        self.sheet_name = ws.title if ws is not None else str(day.day)
        src = ws or sobook.template_sheet(wb, day)
        info = sobook.analyze(src) if src is not None else None
        codes = list(info.class_rows) if info else self.roster.class_codes()
        st = DayState.new(day, self.roster, codes)
        self.template_name = src.title if src is not None else None
        self.layout = info.meal if info else MealLayout()
        if info:
            recs = sobook.read_records(info)
            self._ensure_classes({c: r.gvcn for c, r in recs.items()})
            for code, rec in recs.items():
                if ws is not None:
                    self.sheet_checks[code] = st.load_class_from_sheet(code, rec.si_so, rec.names, rec.note)
                    st.cd(code).keep_sheet_si_so = True
                else:
                    c = st.cd(code)
                    c.sheet_si_so = rec.si_so
                    c.note = rec.note
        st.exists_in_book = ws is not None
        if ws is None and src is not None:
            st.template_date = sobook.sheet_date(src)
        if ws is not None and info:
            for mr in info.meal.rows:  # số vắng ghép đã ghi trong sổ (dùng khi chưa có danh sách ghép đủ)
                if mr.code and mr.m_cell:
                    v = ws[mr.m_cell].value
                    st.manual_m[mr.code] = int(v) if isinstance(v, (int, float)) else 0
        flags = self.data.reported_flags(day)
        if ws is not None and not flags:
            for c in st.classes.values():  # ngày đã ghi sổ trước khi dùng app: coi như các lớp đã báo
                c.reported = True
        for code, flag in flags.items():
            if code in st.classes:
                st.cd(code).reported = flag
        self.state = st
        return st

    # ---------------------------------------------------------------- nhập liệu
    def push_undo(self) -> None:
        if self.state is not None:
            self._undo.append(self.state.snapshot())
            del self._undo[:-50]
        self.dirty = True

    def undo(self) -> bool:
        if not self._undo:
            return False
        self.state = self._undo.pop()
        self.dirty = True
        return True

    def paste(self, text: str) -> list[PasteOutcome]:
        """Đọc tin nhắn dán vào và ghi vào ngày đang mở.

        Tin Zalo có ngày giờ: chỉ lấy tin đúng ngày đang mở (tin ngày khác đếm vào `paste_skipped`),
        theo thứ tự giờ gửi nên tin sau của cùng lớp ghi đè tin trước."""
        assert self.state is not None
        known = set(self.state.classes)
        out = []
        self.paste_skipped = {}
        msgs = []
        for m in parse_messages(text, known):
            if m.when is not None and m.when.date() != self.day:
                d = m.when.date()
                self.paste_skipped[d] = self.paste_skipped.get(d, 0) + 1
            else:
                msgs.append(m)
        if any(m.is_report for m in msgs):
            self.push_undo()
        for msg in msgs:
            po = PasteOutcome(msg)
            if msg.cls is None:
                self._resolve_class(msg, po)
                if msg.cls is None:
                    out.append(po)
                    continue
            elif msg.sender:
                self._learn_sender(msg.sender, msg.cls)
            if msg.schedule:
                self._schedule_suggestions(msg, po)
            if not msg.is_report:
                out.append(po)
                continue
            if msg.full:
                self.state.mark_full(msg.cls)
                c = self.state.cd(msg.cls)
                c.teacher_present, c.teacher_total = msg.present, msg.total
            else:
                po.results = self.state.apply_message(msg)
            c = self.state.cd(msg.cls)
            c.guess = po.class_how if msg.guessed else ""
            c.guess_weak = msg.guessed and po.class_by == "names"
            out.append(po)
        seen = {m.raw for m in self.unplaced}
        self.unplaced += [po.message for po in out
                          if po.message.cls is None and po.message.names and po.message.raw not in seen]
        return out

    def assign_unplaced(self, msg: ParsedMessage, code: str) -> list[matcher.MatchResult]:
        """Giáo vụ chọn lớp cho tin không rõ lớp: ghi tin vào lớp đó."""
        assert self.state is not None
        self.push_undo()
        msg.cls, msg.guessed = code, False
        if msg.full:
            self.state.mark_full(code)
            results = []
        else:
            results = self.state.apply_message(msg)
        self.unplaced = [m for m in self.unplaced if m is not msg]
        return results

    def dismiss_unplaced(self, msg: ParsedMessage) -> None:
        self.unplaced = [m for m in self.unplaced if m is not msg]

    # ---------------------------------------------------------------- việc còn dở trước khi lưu
    def pending_names(self) -> dict[str, list[str]]:
        """Tên GV ghi mà app chưa xác định được em nào (chưa tính vắng), theo lớp."""
        assert self.state is not None
        out: dict[str, list[str]] = {}
        for code, res in self.state.pending:
            out.setdefault(code, []).append(res.query)
        return out

    def unresolved_count(self) -> int:
        """Số tên GV báo vắng chưa xác định được em nào (chưa tính vắng). Còn > 0 thì chưa cho lưu."""
        if self.state is None:
            return 0
        return len(self.state.pending) + sum(len(m.names) for m in self.unplaced)

    def guessed_classes(self) -> dict[str, str]:
        """Lớp có tin do app tự đoán lớp, giáo vụ chưa xem lại: mã lớp -> cách đoán."""
        assert self.state is not None
        return {code: c.guess for code, c in self.state.classes.items() if c.guess}

    def confirm_guess(self, code: str) -> None:
        assert self.state is not None
        self.push_undo()
        c = self.state.cd(code)
        c.guess, c.guess_weak = "", False

    def _resolve_class(self, msg: ParsedMessage, po: PasteOutcome) -> None:
        """Tìm lớp cho tin không ghi mã lớp: theo tên Zalo người gửi, rồi theo tên học sinh."""
        from .core.senders import match_teacher, sender_key

        assert self.state is not None
        known = set(self.state.classes)
        by_sender, how = None, ""
        if msg.sender_class in known:
            by_sender, how = msg.sender_class, f"theo mã lớp trong tên Zalo {msg.sender}"
        elif msg.sender:
            gvcn = {c: info.gvcn for c, info in self.roster.classes.items() if c in known and info.gvcn}
            by_sender, how = match_teacher(msg.sender, gvcn), f"theo tên Zalo {msg.sender} giống GVCN"
            if by_sender is None:
                learned = self.settings.get(sender_key(msg.sender))
                by_sender = learned if learned in known else None
                how = f"theo lớp {msg.sender} đã báo lần trước"
        guesses = matcher.guess_class(self.roster, msg.names, self.day) if msg.names else []
        po.guessed_from = guesses
        if by_sender and (not msg.names or by_sender in guesses or len(guesses) != 1):
            msg.cls, msg.guessed, po.class_by = by_sender, True, "sender"
            po.class_how = how
        elif len(guesses) == 1:
            msg.cls, msg.guessed, po.class_by = guesses[0], True, "names"
            po.class_how = "theo tên học sinh trong tin"
        if msg.cls and msg.sender:
            self._learn_sender(msg.sender, msg.cls)

    def _learn_sender(self, sender: str, code: str) -> None:
        from .core.senders import sender_key

        key = sender_key(sender)
        if key != "zalo|" and self.settings.get(key) != code:
            self.settings[key] = code
            self.data_dirty = True

    def _schedule_suggestions(self, msg: ParsedMessage, po: PasteOutcome) -> None:
        from .core.notes import suggest_from_note

        if not msg.cls:
            return
        seen = msg.when.date() if msg.when else self.day
        po.suggestions = suggest_from_note(self.roster, msg.cls, _relative_dates(msg.schedule, seen), seen)
        po.schedule_known = not po.suggestions

    def apply_paste_suggestion(self, sug) -> str:
        """Áp dụng 1 đề nghị từ tin báo lịch; thêm ghi chú cột H như giáo vụ vẫn viết. Trả về ghi chú."""
        from .core import schedule as sch
        from .core.notes import apply_suggestions

        if not apply_suggestions(self.roster, [sug]):
            return ""
        self.data_dirty = True
        s = sug.student or next((x for x in reversed(self.roster.students) if x.ho_ten == sug.new_name), None)
        ch = sug.changes
        if ch.get("cat_tu"):
            note = sch.note_text(sch.CUT, [s], ch["cat_tu"])
        elif ch.get("cat_thu6"):
            note = sch.note_text(sch.FRIDAY, [s], None)
        elif ch.get("ban_tru"):
            note = sch.note_text(sch.ADD, [s], ch.get("bt_tu") or self.day)
        else:
            note = ""
        if note and self.state is not None and sug.code in self.state.classes:
            self.push_undo()
            c = self.state.cd(sug.code)
            if note not in c.note:
                c.note = f"{c.note}, {note}" if c.note else note
        self.save_data()
        return note

    # ---------------------------------------------------------------- lưu & in
    def si_so_conflicts(self) -> list[tuple[str, int, int]]:
        """Lớp có sĩ số theo danh sách khác sĩ số đang ghi trong sổ."""
        assert self.state is not None
        out = []
        for code, c in self.state.classes.items():
            if self.state.si_so_mismatch(code):
                out.append((code, self.roster.si_so(code, self.day), c.sheet_si_so))
        return out

    def keep_sheet_si_so(self, codes: list[str]) -> None:
        assert self.state is not None
        for code in codes:
            self.state.cd(code).keep_sheet_si_so = True

    def meals(self) -> MealSummary:
        assert self.state is not None
        return self.state.meals(self.layout)

    def book_locked(self) -> bool:
        return bool(self.book_path and is_locked(self.book_path))

    def save(self) -> str:
        """Ghi ngày đang mở vào sổ (đọc lại file mới nhất trước khi ghi). Trả về tên sheet."""
        assert self.state is not None
        if not self.book_path:
            raise FileNotFoundError("Chưa chọn sổ điểm danh tháng.")
        wb = sobook.load_book(self.book_path)
        sheets = sobook.day_sheets(wb)
        ws = sheets.get(self.day.day)
        if ws is None:
            tpl = sobook.template_sheet(wb, self.day)
            if tpl is None:
                raise ValueError("Sổ không có sheet điểm danh nào để làm mẫu.")
            ws = sobook.create_day_sheet(wb, self.day, tpl)
        info = sobook.analyze(ws)
        if info is None:
            raise ValueError(f"Sheet {ws.title} không đúng mẫu điểm danh.")
        summary = self.state.meals(info.meal)
        sobook.write_day(info, self.state, summary)
        sobook.select_only(wb, ws)
        sobook.save_book(wb, self.book_path)
        self.sheet_name = ws.title
        self.state.exists_in_book = True
        for code in self.state.classes:
            c = self.state.cd(code)
            c.sheet_si_so = self.state.si_so(code)  # lấy sĩ số vừa ghi TRƯỚC khi chuyển sang "giữ sĩ số trong sổ"
            c.keep_sheet_si_so = True
            self.data.set_log(self.day, code, self.state.cd(code).reported, self.state.si_so(code), self.state.absent_names(code))
        self.dirty = False
        self.warning = ""
        try:
            self.save_data()
        except Exception as e:  # noqa: BLE001 - sổ đã lưu xong; chỉ nhật ký của app chưa ghi được
            self.data_dirty = True
            self.warning = f"Sổ đã lưu, nhưng chưa ghi được file dữ liệu của app ({e}). Hãy đóng file đó nếu đang mở."
        return ws.title

    def print_areas(self) -> list[str]:
        info = None
        if self.book_path:
            wb = self._book(self.book_path)
            ws = sobook.day_sheets(wb).get(self.day.day)
            info = sobook.analyze(ws) if ws is not None else None
        return [a for a in (self.page1_area(info), self.page2_area()) if a]

    # ---------------------------------------------------------------- lịch thêm / cắt
    def schedule(self, kind: str, code: str, ids: list[str], start: dt.date | None, end: dt.date | None = None,
                 new_name: str = "", new_dob: str = "", add_note: bool = True) -> str:
        """Đặt lịch cho các em (hoặc thêm em mới chưa có trong danh sách). Trả về ghi chú đã thêm vào cột H."""
        from .core import schedule as sch
        from .core.models import Student

        students = [s for s in self.roster.students if s.id in ids]
        if new_name:
            s = Student(self.roster.next_id(), code, new_name, new_dob, ban_tru=False, ghi_chu="thêm bằng app")
            self.roster.students.append(s)
            students.append(s)
        if not students:
            return ""
        sch.apply(self.roster, kind, students, start, end)
        self.data_dirty = True
        note = sch.note_text(kind, students, start, end) if add_note else ""
        if note and self.state is not None and code in self.state.classes:
            self.push_undo()
            c = self.state.cd(code)
            if note not in c.note:
                c.note = f"{c.note}, {note}" if c.note else note
        self.save_data()
        return note

    def cancel_schedule(self, item) -> None:
        from .core import schedule as sch

        sch.cancel(self.roster, item)
        self.data_dirty = True
        self.save_data()

    def schedule_items(self):
        from .core import schedule as sch

        return sch.items(self.roster, self.day)

    # ---------------------------------------------------------------- ghi chú trong sổ
    def note_suggestions(self):
        """Đọc ghi chú cột H ở mọi sheet của sổ, gợi ý cập nhật danh sách (cắt thứ 6, cắt/thêm theo ngày…)."""
        from .core.notes import suggest_from_book

        if not self.book_path:
            return []
        wb = self._book(self.book_path)
        notes = []
        for _, ws in sorted(sobook.day_sheets(wb).items()):
            info = sobook.analyze(ws)
            d = sobook.sheet_date(ws)
            if not info or not d:
                continue
            for code, rec in sobook.read_records(info).items():
                if rec.note:
                    notes.append((d, code, rec.note))
                # ghi chú lẫn trong ô tên vắng ("cắt …", "thêm …")
                from .core.day import split_sheet_cell

                for extra in split_sheet_cell(rec.names)[1]:
                    notes.append((d, code, extra))
        return suggest_from_book(self.roster, notes)

    def apply_note_suggestions(self, sugs) -> int:
        from .core.notes import apply_suggestions

        n = apply_suggestions(self.roster, sugs)
        if n:
            self.data_dirty = True
            self.save_data()
        return n

    # ---------------------------------------------------------------- ghép lớp
    def ghep_students(self, code: str):
        return self.roster.sort_like_roster(self.roster.boarding(code, self.day))

    def assign_ghep(self, ids: list[str], group: str) -> None:
        for s in self.roster.students:
            if s.id in ids:
                s.an_ghep = group
        self.data_dirty = True

    def assign_ghep_text(self, cls: str, text: str) -> tuple[int, list[str]]:
        """Dán danh sách ghép, mỗi dòng 'mã nhóm: tên, tên…' (vd. '1.2: Bảo Ngọc, Gia Hân')."""
        import re

        from .core.parser import extract_names

        n, problems = 0, []
        for line in text.splitlines():
            m = re.match(r"\s*(\d)\s*[./]\s*(\d)\s*[:\-–]?\s*(.*)$", line)
            if not m:
                if line.strip():
                    problems.append(f"Không hiểu dòng: {line.strip()}")
                continue
            group = f"{m.group(1)}.{m.group(2)}"
            for q in extract_names(m.group(3)):
                r = matcher.match_name(self.roster, cls, q, self.day)
                if r.status == matcher.OK and r.student:
                    r.student.an_ghep = group
                    self.data_dirty = True
                    n += 1
                else:
                    problems.append(f"{group}: '{q}' {'trùng nhiều em' if r.status == matcher.AMBIGUOUS else 'không tìm thấy'}")
        return n, problems


def _relative_dates(text: str, seen: dt.date) -> str:
    """Đổi "từ mai", "tháng sau" trong tin báo lịch thành ngày cụ thể (tính từ ngày gửi tin)."""
    import re

    from .core.schedule import next_month_start

    nxt = next_school_day(seen)
    text = re.sub(r"(?i)từ\s*(?:ngày\s*)?mai", f"từ {nxt.day}/{nxt.month}", text)
    text = re.sub(r"(?i)(?:đầu\s*)?tháng\s*sau", f"tháng {next_month_start(seen).month}", text)
    return text


def _col(i: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(i)


def search_students(roster, text: str, cls: str | None = None, only_bt: bool = False):
    f = fold(text)
    out = []
    for s in roster.students:
        if cls and s.lop != cls and s.lop_tinh != cls:
            continue
        if only_bt and not s.ban_tru:
            continue
        if f and f not in fold(s.ho_ten):
            continue
        out.append(s)
    return out
