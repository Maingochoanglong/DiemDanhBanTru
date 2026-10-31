"""Đọc danh sách học sinh (.xls hoặc .xlsx).

Hai dạng file:
- file CDHS của trường: em tô chữ đỏ là ăn bán trú;
- file danh sách bán trú gọn: cột "Ăn bán trú" đánh X, kèm cột "Ghép lớp" (vd. 1.2).
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from ..core.models import Roster, Student
from ..core.text import clean_spaces, fmt_date, fold, norm_class, parse_date


@dataclass
class RosterRow:
    lop: str
    ho_ten: str
    ngay_sinh: str
    red: bool  # ăn bán trú (chữ đỏ hoặc X)
    phan_loai: str = ""
    ghep: str | None = None  # None = file không có cột "Ghép lớp"

    @property
    def key(self) -> tuple[str, str]:
        return fold(self.ho_ten), self.ngay_sinh


class RosterFormatError(Exception):
    pass


def _is_red_rgb(r: int, g: int, b: int) -> bool:
    return r >= 180 and g <= 100 and b <= 100


def _find_columns(header: list[str]) -> dict[str, int] | None:
    cols: dict[str, int] = {}
    for i, h in enumerate(header):
        f = fold(str(h or ""))
        if "lop" not in cols and ("ten lop" in f or f in ("lop", "lop hoc")):
            cols["lop"] = i
        elif f in ("ho ten", "ho va ten", "ho ten hoc sinh", "ho va ten hoc sinh") and "ten" not in cols:
            cols["ten"] = i
        elif f.startswith("ngay sinh") and "ns" not in cols:
            cols["ns"] = i
        elif f.startswith("phan loai") and "pl" not in cols:
            cols["pl"] = i
        elif (f.startswith(("an ban tru", "ban tru")) or f in ("bt", "an bt")) and "bt" not in cols:
            # cột đánh X; không nhận nhầm cột "Loại bán trú" (ghi "Thường") của file CDHS
            cols["bt"] = i
        elif "ghep" in f and "ghep" not in cols:
            cols["ghep"] = i
    if "lop" in cols and "ten" in cols:
        return cols
    return None


def _is_x(v) -> bool:
    return fold(str(v or "")) in ("x", "co", "1", "v", "ban tru", "true")


def _fmt_ghep(v) -> str:
    t = clean_spaces(str(v if v is not None else "")).replace(",", ".")
    if isinstance(v, float):
        t = f"{v:.1f}"
    return t


def _fmt_dob(v) -> str:
    d = parse_date(v)
    return fmt_date(d) if d else clean_spaces(str(v or ""))


def _read_xls(path: Path) -> list[RosterRow]:
    import xlrd

    book = xlrd.open_workbook(str(path), formatting_info=True)
    for sh in book.sheets():
        for hr in range(min(10, sh.nrows)):
            cols = _find_columns([str(v) for v in sh.row_values(hr)])
            if not cols:
                continue
            rows = []
            for r in range(hr + 1, sh.nrows):
                name = clean_spaces(str(sh.cell_value(r, cols["ten"])))
                lop = norm_class(str(sh.cell_value(r, cols["lop"])))
                if not name or not lop:
                    continue
                font_idx = book.xf_list[sh.cell_xf_index(r, cols["ten"])].font_index
                runs = sh.rich_text_runlist_map.get((r, cols["ten"]))
                if runs and runs[0][0] == 0:
                    font_idx = runs[0][1]
                rgb = book.colour_map.get(book.font_list[font_idx].colour_index)
                red = bool(rgb) and _is_red_rgb(*rgb)
                dob = ""
                if "ns" in cols:
                    cell = sh.cell(r, cols["ns"])
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        dob = fmt_date(xlrd.xldate_as_datetime(cell.value, book.datemode).date())
                    else:
                        dob = _fmt_dob(cell.value)
                pl = clean_spaces(str(sh.cell_value(r, cols["pl"]))) if "pl" in cols else ""
                if "bt" in cols:
                    red = _is_x(sh.cell_value(r, cols["bt"]))
                ghep = _fmt_ghep(sh.cell_value(r, cols["ghep"])) if "ghep" in cols else None
                rows.append(RosterRow(lop, name, dob, red, pl, ghep))
            return rows
    raise RosterFormatError("Không tìm thấy cột 'Tên lớp học' và 'Họ tên' trong file.")


def _openpyxl_red(cell) -> bool:
    from openpyxl.styles.colors import COLOR_INDEX

    color = cell.font.color if cell.font else None
    if color is None:
        return False
    hexv = None
    if color.type == "rgb" and isinstance(color.rgb, str):
        hexv = color.rgb[-6:]
    elif color.type == "indexed" and color.indexed is not None and color.indexed < len(COLOR_INDEX):
        hexv = COLOR_INDEX[color.indexed][-6:]
    if not hexv:
        return False
    try:
        return _is_red_rgb(int(hexv[0:2], 16), int(hexv[2:4], 16), int(hexv[4:6], 16))
    except ValueError:
        return False


def _read_xlsx(path: Path) -> list[RosterRow]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    for ws in wb.worksheets:
        for hr in range(1, min(10, ws.max_row) + 1):
            header = [c.value for c in ws[hr]]
            cols = _find_columns([str(v) if v is not None else "" for v in header])
            if not cols:
                continue
            rows = []
            for r in range(hr + 1, ws.max_row + 1):
                name_cell = ws.cell(r, cols["ten"] + 1)
                name = clean_spaces(str(name_cell.value or ""))
                lop = norm_class(str(ws.cell(r, cols["lop"] + 1).value or ""))
                if not name or not lop:
                    continue
                dob = ""
                if "ns" in cols:
                    dob = _fmt_dob(ws.cell(r, cols["ns"] + 1).value)
                pl = clean_spaces(str(ws.cell(r, cols["pl"] + 1).value or "")) if "pl" in cols else ""
                red = _is_x(ws.cell(r, cols["bt"] + 1).value) if "bt" in cols else _openpyxl_red(name_cell)
                ghep = _fmt_ghep(ws.cell(r, cols["ghep"] + 1).value) if "ghep" in cols else None
                rows.append(RosterRow(lop, name, dob, red, pl, ghep))
            return rows
    raise RosterFormatError("Không tìm thấy cột 'Tên lớp học' và 'Họ tên' trong file.")


def read_cdhs(path: str | Path) -> list[RosterRow]:
    """Đọc theo nội dung file (file .xls cũ đặt đuôi .xlsx vẫn đọc được)."""
    path = Path(path)
    with open(path, "rb") as f:
        head = f.read(8)
    if head.startswith(b"\xd0\xcf\x11\xe0"):  # định dạng Excel 97-2003 (.xls)
        return _read_xls(path)
    return _read_xlsx(path)


def flag_mismatches(rows: list[RosterRow]) -> list[RosterRow]:
    """Em có màu chữ đỏ không khớp cột 'Phân loại HS'."""
    out = []
    for r in rows:
        if not r.phan_loai:
            continue
        bt = fold(r.phan_loai).startswith("ban tru")
        if bt != r.red:
            out.append(r)
    return out


# ---------------- so sánh với danh sách đang có ----------------
@dataclass
class Change:
    kind: str  # new | removed | class | boarding | ghep
    text: str
    row: RosterRow | None = None
    student: Student | None = None
    apply: bool = True


def diff_roster(roster: Roster, rows: list[RosterRow]) -> list[Change]:
    existing = {s.key(): s for s in roster.students}
    seen = set()
    changes: list[Change] = []
    for row in rows:
        s = existing.get(row.key)
        seen.add(row.key)
        if s is None:
            bt = "ăn BT" if row.red else "không ăn BT"
            changes.append(Change("new", f"Thêm HS mới: {row.ho_ten} ({row.lop}, {bt})", row))
            continue
        if s.lop != row.lop:
            changes.append(Change("class", f"Đổi lớp: {s.ho_ten} {s.lop} → {row.lop}", row, s))
        if s.ban_tru != row.red and not (not row.red and s.cat_tu):  # em đã có lịch cắt thì giữ nguyên
            what = "bắt đầu ăn BT" if row.red else "thôi ăn BT"
            changes.append(Change("boarding", f"{s.ho_ten} ({row.lop}) {what}", row, s))
        if row.ghep is not None and row.ghep != s.an_ghep:
            what = f"ăn ghép {row.ghep}" if row.ghep else "bỏ ăn ghép"
            changes.append(Change("ghep", f"{s.ho_ten} ({row.lop}) {what}", row, s))
    for key, s in existing.items():
        if key not in seen:
            changes.append(Change("removed", f"Không còn trong danh sách: {s.ho_ten} ({s.lop})", None, s, apply=False))
    return changes


def apply_changes(roster: Roster, changes: list[Change]) -> int:
    n = 0
    for c in changes:
        if not c.apply:
            continue
        if c.kind == "new" and c.row:
            roster.students.append(
                Student(roster.next_id(), c.row.lop, c.row.ho_ten, c.row.ngay_sinh, c.row.red, an_ghep=c.row.ghep or "")
            )
        elif c.kind == "class" and c.student and c.row:
            c.student.lop = c.row.lop
        elif c.kind == "boarding" and c.student and c.row:
            c.student.ban_tru = c.row.red
        elif c.kind == "ghep" and c.student and c.row:
            c.student.an_ghep = c.row.ghep or ""
        elif c.kind == "removed" and c.student:
            roster.students.remove(c.student)
        n += 1
    roster.invalidate()
    return n
