"""Sổ điểm danh tháng: mỗi ngày 1 sheet (tên = số ngày), mỗi tháng 1 file.

Không cố định địa chỉ ô: tìm hàng tiêu đề "TÊN HỌC SINH VẮNG", tìm hàng lớp theo mã lớp ở
cột LỚP, tìm bảng suất ăn theo công thức. App chỉ ghi giá trị vào các ô số liệu; định dạng,
công thức, chữ ký giữ nguyên như sheet mẫu.
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..core.day import DayState
from ..core.meals import MealLayout, MealRow, MealSummary
from ..core.text import clean_spaces, fmt_date, fold, norm_class
from .files import safe_save

_REF = re.compile(r"(?<![A-Z:$])\$?([A-Z]{1,3})\$?(\d+)(?!\s*:)(?![\d(])")
_TITLE_DATE = re.compile(r"(\d{1,2})\s*/\s*(\d{1,2})\s*/\s*(\d{4})")


@dataclass
class SheetInfo:
    ws: Worksheet
    header_row: int
    cols: dict[str, int]  # lop, gvcn, siso, vang, hien, ten, ghichu (chỉ số cột 1-based)
    class_rows: dict[str, int]
    total_row: int
    title_cell: str | None = None
    sign_cell: str | None = None
    meal: MealLayout = field(default_factory=MealLayout)

    def cell(self, code: str, key: str):
        return self.ws.cell(self.class_rows[code], self.cols[key])


@dataclass
class ClassRecord:
    code: str
    gvcn: str
    si_so: int | None
    vang: int | None
    names: str
    note: str


# ---------------------------------------------------------------- nhận diện sheet
def analyze(ws: Worksheet) -> SheetInfo | None:
    header_row = None
    cols: dict[str, int] = {}
    for r in range(1, min(ws.max_row, 20) + 1):
        for c in range(1, min(ws.max_column, 30) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and "ten hoc sinh vang" in fold(v):
                header_row, cols["ten"] = r, c
                break
        if header_row:
            break
    if not header_row:
        return None
    for c in range(1, cols["ten"]):
        f = fold(str(ws.cell(header_row, c).value or ""))
        if f == "lop" or f.startswith("lop"):
            cols.setdefault("lop", c)
        elif "giao vien" in f or "gvcn" in f:
            cols.setdefault("gvcn", c)
        elif f.startswith("si so") or f.startswith("sy so"):
            cols.setdefault("siso", c)
        elif f.startswith("vang"):
            cols.setdefault("vang", c)
        elif "hien dien" in f:
            cols.setdefault("hien", c)
    if not {"lop", "siso", "vang"} <= set(cols):
        return None
    cols.setdefault("hien", cols["vang"] + 1)
    cols["ghichu"] = cols["ten"] + 1

    class_rows: dict[str, int] = {}
    total_row = header_row
    for r in range(header_row + 1, ws.max_row + 1):
        v = ws.cell(r, cols["lop"]).value
        code = norm_class(str(v)) if isinstance(v, str) else ""
        if code and code not in class_rows:
            class_rows[code] = r
            total_row = max(total_row, r)
        a = ws.cell(r, 1).value
        if isinstance(a, str) and "tong cong" in fold(a):
            total_row = r
            break
    if not class_rows:
        return None

    info = SheetInfo(ws, header_row, cols, class_rows, total_row)
    for r in range(1, header_row):
        for c in range(1, min(ws.max_column, 30) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and "ngay" in fold(v) and _TITLE_DATE.search(v):
                info.title_cell = ws.cell(r, c).coordinate
    for r in range(total_row + 1, min(ws.max_row, total_row + 6) + 1):
        for c in range(1, min(ws.max_column, 30) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and "thang" in fold(v) and "nam" in fold(v):
                info.sign_cell = ws.cell(r, c).coordinate
                break
        if info.sign_cell:
            break
    info.meal = analyze_meals(ws, info)
    return info


def _refs(formula: str) -> list[tuple[str, int]]:
    return [(m.group(1), int(m.group(2))) for m in _REF.finditer(formula.upper())]


def analyze_meals(ws: Worksheet, info: SheetInfo) -> MealLayout:
    """Đọc bảng suất ăn bên dưới: nhóm nào cộng những lớp nào, hàng nào nhận HS ghép."""
    hien = get_column_letter(info.cols["hien"])
    row_to_class = {r: c for c, r in info.class_rows.items()}
    formula_cells = []
    for row in ws.iter_rows(min_row=info.total_row + 1, max_row=ws.max_row):
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                refs = _refs(v)
                classes = [row_to_class[r] for col, r in refs if col == hien and r in row_to_class]
                others = [(col, r) for col, r in refs if not (col == hien and r in row_to_class)]
                formula_cells.append((cell, classes, others))
    if not formula_cells:
        return MealLayout()

    def label_of(cell) -> str:
        nxt = ws.cell(cell.row, cell.column + 1).value
        return clean_spaces(nxt) if isinstance(nxt, str) else ""

    score: dict[int, int] = {}
    for cell, classes, others in formula_cells:
        if label_of(cell) and (classes or others):
            score[cell.column] = score.get(cell.column, 0) + 1
    if not score:
        return MealLayout()
    meal_col = max(score, key=lambda k: score[k])

    layout = MealLayout()
    used_rows: list[int] = []
    referenced: set[str] = set()
    for cell, classes, others in formula_cells:
        referenced.update(classes)
        if cell.column == meal_col and label_of(cell):
            mr = MealRow(cell.row, label_of(cell), cell.coordinate, classes)
            for col, r in others:
                sub = ws[f"{col}{r}"].value
                m = re.match(r"=\s*\$?([A-Z]{1,3})\$?(\d+)\s*-\s*\$?([A-Z]{1,3})\$?(\d+)\s*$", str(sub or "").upper())
                if m:
                    mr.o_cell, mr.m_cell = f"{m.group(1)}{m.group(2)}", f"{m.group(3)}{m.group(4)}"
                    break
            layout.rows.append(mr)
            used_rows.append(cell.row)
        elif classes:
            layout.self_classes.extend(c for c in classes if c not in layout.self_classes)
            used_rows.append(cell.row)
    if not layout.rows:
        return MealLayout()
    if any(r.code for r in layout.rows):
        layout.ghep_classes = [c for c in info.class_rows if c not in referenced]
    for mr in layout.rows:
        if mr.code and mr.o_cell:
            v = ws[mr.o_cell].value
            layout.o_values[mr.code] = int(v) if isinstance(v, (int, float)) else 0

    # Vùng in: khung bao các ô có nội dung từ hàng tiêu đề "Suất ăn" tới hàng cuối của bảng
    r1, r2 = min(used_rows) - 1, max(used_rows)
    cmin, cmax = 10**6, 0
    for row in ws.iter_rows(min_row=r1, max_row=r2):
        for cell in row:
            if cell.value not in (None, ""):
                cmin, cmax = min(cmin, cell.column), max(cmax, cell.column)
    if cmax:
        layout.area = f"{get_column_letter(cmin)}{r1}:{get_column_letter(cmax)}{r2}"
    return layout


# ---------------------------------------------------------------- đọc
def day_sheets(wb) -> dict[int, Worksheet]:
    out = {}
    for ws in wb.worksheets:
        t = ws.title.strip()
        if t.isdigit() and 1 <= int(t) <= 31:
            out[int(t)] = ws
    return out


def sheet_date(ws: Worksheet) -> dt.date | None:
    info = analyze(ws)
    if info and info.title_cell:
        m = _TITLE_DATE.search(str(ws[info.title_cell].value))
        if m:
            try:
                return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
            except ValueError:
                return None
    return None


def book_month(wb, path: str | Path | None = None) -> tuple[int, int] | None:
    """(tháng, năm) của sổ, đọc từ tiêu đề sheet ngày; dự phòng: tên file 'T10'."""
    for _, ws in sorted(day_sheets(wb).items(), reverse=True):
        d = sheet_date(ws)
        if d:
            return d.month, d.year
    if path:
        m = re.search(r"T(\d{1,2})(?!\d)", Path(path).stem, re.IGNORECASE)
        if m:
            return int(m.group(1)), dt.date.today().year
    return None


def read_records(info: SheetInfo) -> dict[str, ClassRecord]:
    out = {}
    for code, r in info.class_rows.items():
        vals = {k: info.ws.cell(r, c).value for k, c in info.cols.items()}
        si, vang = vals.get("siso"), vals.get("vang")
        out[code] = ClassRecord(
            code,
            clean_spaces(str(vals.get("gvcn") or "")),
            int(si) if isinstance(si, (int, float)) else None,
            int(vang) if isinstance(vang, (int, float)) else None,
            clean_spaces(str(vals.get("ten") or "")),
            clean_spaces(str(vals.get("ghichu") or "")),
        )
    return out


def template_sheet(wb, day: dt.date) -> Worksheet | None:
    """Sheet gần nhất trước ngày `day` (hoặc sheet cuối cùng) để làm mẫu."""
    sheets = day_sheets(wb)
    before = [d for d in sheets if d < day.day]
    pick = max(before) if before else (max(sheets) if sheets else None)
    if pick is not None:
        return sheets[pick]
    for ws in reversed(wb.worksheets):
        if analyze(ws):
            return ws
    return None


# ---------------------------------------------------------------- ghi
def create_day_sheet(wb, day: dt.date, template: Worksheet) -> Worksheet:
    """Sao chép sheet mẫu thành sheet ngày mới, xóa số liệu cũ (giữ ghi chú, chữ ký)."""
    new = wb.copy_worksheet(template)
    new.title = str(day.day)
    new.sheet_view.zoomScale = template.sheet_view.zoomScale
    if template.freeze_panes:
        new.freeze_panes = template.freeze_panes
    # đặt đúng thứ tự ngày
    wb._sheets.remove(new)
    sheets = day_sheets(wb)
    idx = len(wb._sheets)
    after = [d for d in sheets if d > day.day]
    if after:
        idx = wb._sheets.index(sheets[min(after)])
    wb._sheets.insert(idx, new)

    info = analyze(new)
    if info:
        for code in info.class_rows:
            info.cell(code, "ten").value = None
            info.cell(code, "vang").value = 0
        for mr in info.meal.rows:
            if mr.m_cell:
                new[mr.m_cell].value = None
        set_dates(info, day)
    return new


def set_dates(info: SheetInfo, day: dt.date) -> None:
    ws = info.ws
    if info.title_cell:
        ws[info.title_cell].value = _TITLE_DATE.sub(fmt_date(day), str(ws[info.title_cell].value), count=1)
    if info.sign_cell:
        v = str(ws[info.sign_cell].value)
        v = re.sub(r"(tháng\s+)\d{1,2}", rf"\g<1>{day.month}", v, count=1)
        v = re.sub(r"(năm\s+)\d{4}", rf"\g<1>{day.year}", v, count=1)
        ws[info.sign_cell].value = v


def write_day(info: SheetInfo, state: DayState, summary: MealSummary | None) -> None:
    for code in info.class_rows:
        if code not in state.classes:
            continue
        names = state.absent_names(code)
        info.cell(code, "siso").value = state.si_so(code)
        info.cell(code, "vang").value = len(names)
        info.cell(code, "ten").value = ", ".join(names) if names else None
        note = state.cd(code).note
        info.cell(code, "ghichu").value = note or None
    if summary:
        for line in summary.lines:
            if line.row.o_cell and line.o is not None:
                info.ws[line.row.o_cell].value = line.o
            if line.row.m_cell:
                info.ws[line.row.m_cell].value = line.m or None


def select_only(wb, ws: Worksheet) -> None:
    """Chỉ chọn 1 sheet khi mở file (tránh Excel mở ở chế độ nhóm nhiều sheet)."""
    for s in wb.worksheets:
        s.sheet_view.tabSelected = s is ws
    wb.active = wb.worksheets.index(ws)


def load_book(path: str | Path):
    return openpyxl.load_workbook(str(path))


def save_book(wb, path: str | Path, backup_dir: str | Path | None = None) -> None:
    safe_save(wb, path, backup_dir)


def month_file_name(prev: Path, month: int) -> Path:
    stem = prev.stem
    if re.search(r"T\d{1,2}(?!\d)", stem):
        idx = [m for m in re.finditer(r"T\d{1,2}(?!\d)", stem)][-1]
        stem = stem[: idx.start()] + f"T{month}" + stem[idx.end():]
    else:
        stem = f"{stem} T{month}"
    return prev.with_name(stem + prev.suffix)


def new_month_book(prev_path: str | Path, day: dt.date, new_path: str | Path) -> None:
    """Tạo sổ tháng mới từ sheet cuối của sổ tháng trước."""
    wb = load_book(prev_path)
    sheets = day_sheets(wb)
    tpl = sheets[max(sheets)] if sheets else template_sheet(wb, day)
    if tpl is None:
        raise ValueError("Sổ tháng trước không có sheet điểm danh nào để làm mẫu.")
    for ws in list(wb.worksheets):
        if ws is not tpl:
            wb.remove(ws)
    tpl.title = str(day.day)
    info = analyze(tpl)
    if info:
        for code in info.class_rows:
            info.cell(code, "ten").value = None
            info.cell(code, "vang").value = 0
        for mr in info.meal.rows:
            if mr.m_cell:
                tpl[mr.m_cell].value = None
        set_dates(info, day)
    select_only(wb, tpl)
    safe_save(wb, new_path)


def find_month_file(folder: str | Path, month: int, year: int | None = None) -> Path | None:
    folder = Path(folder)
    if not folder.is_dir():
        return None
    cands = []
    for p in folder.glob("*.xlsx"):
        if p.name.startswith(("~$", ".")):
            continue
        if re.search(rf"T0?{month}(?!\d)", p.stem, re.IGNORECASE):
            cands.append(p)
    if not cands:
        return None
    if year and len(cands) > 1:
        for p in cands:
            try:
                wb = openpyxl.load_workbook(str(p), read_only=False)
                bm = book_month(wb, p)
                if bm and bm == (month, year):
                    return p
            except Exception:  # noqa: BLE001 - file hỏng thì bỏ qua
                continue
    return sorted(cands, key=lambda p: p.stat().st_mtime)[-1]
