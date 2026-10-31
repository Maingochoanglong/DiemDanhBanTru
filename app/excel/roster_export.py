"""Xuất danh sách học sinh ăn bán trú dạng gọn: Tên lớp | Họ tên | Ngày sinh | Ăn bán trú (X) | Ghép lớp.

File xuất ra nhập lại được vào app (cột X thay cho tô chữ đỏ).
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..core.models import Roster
from ..core.text import class_sort_key, fmt_date
from .files import safe_save

HEADERS = [("Tên lớp học", 12), ("Họ tên", 32), ("Ngày sinh", 13), ("Ăn bán trú", 12), ("Ghép lớp", 11)]
FONT = "Times New Roman"
CHANGED_FILL = PatternFill("solid", fgColor="FFFFF2CC")
HEADER_FILL = PatternFill("solid", fgColor="FFDDEBF7")
THIN = Side(style="thin", color="FF999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def export_roster(
    roster: Roster,
    path: str | Path,
    as_of: dt.date,
    highlight: set[str] | None = None,
    changes: list[list] | None = None,
    backup: bool = True,
) -> int:
    """Ghi danh sách tính theo ngày `as_of` (X = ăn bán trú ngày đó). Trả về số em đánh X.

    highlight: id các em có thay đổi (tô vàng); changes: các dòng cho sheet "Thay đổi"."""
    highlight = highlight or set()
    pos = {s.id: i for i, s in enumerate(roster.students)}
    students = sorted(roster.students, key=lambda s: (class_sort_key(s.lop), pos[s.id]))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DanhSach"
    _header(ws, HEADERS)
    n_x = 0
    for s in students:
        eats = s.is_boarding(as_of)
        n_x += eats
        ws.append([f"Lớp {s.lop}", s.ho_ten, s.ngay_sinh, "X" if eats else None, s.an_ghep or None])
        r = ws.max_row
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name=FONT, size=12, bold=(c == 4 and eats))
            cell.border = BORDER
            if c in (1, 3, 4, 5):
                cell.alignment = Alignment(horizontal="center")
            if s.id in highlight:
                cell.fill = CHANGED_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9
    ws.print_title_rows = "1:1"

    if changes:
        wc = wb.create_sheet("Thay đổi")
        _header(wc, [("Lớp", 8), ("Họ tên", 30), ("Thay đổi", 48), ("Theo ghi chú trong sổ", 60)])
        for row in changes:
            wc.append(row)
            for c in range(1, 5):
                cell = wc.cell(wc.max_row, c)
                cell.font = Font(name=FONT, size=12)
                cell.border = BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    info = wb.create_sheet("Ghi chú")
    info.append([f"Danh sách tính theo ngày {fmt_date(as_of)}. X = ăn bán trú; trống = không ăn."])
    info.append(["Ghép lớp: nhóm ăn chung của em lớp ghép (vd. 1.2 = ăn cùng lớp 1.2)."])
    info.append(["Dòng tô vàng: có thay đổi so với danh sách trước (xem sheet \"Thay đổi\")."])
    info.column_dimensions["A"].width = 90

    if backup:
        safe_save(wb, path)
    else:
        wb.save(str(path))
    return n_x


def _header(ws, cols) -> None:
    ws.append([t for t, _ in cols])
    for i, (_, w) in enumerate(cols, 1):
        cell = ws.cell(1, i)
        cell.font = Font(name=FONT, size=12, bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
