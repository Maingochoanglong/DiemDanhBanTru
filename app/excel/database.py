"""File dữ liệu của app: DuLieu_BanTru.xlsx (mở được bằng Excel để xem/sửa)."""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..core.models import ClassInfo, Roster, Student
from ..core.text import class_sort_key, clean_spaces, fmt_date, fold, norm_class, parse_date
from .files import safe_save

HS_COLS = [
    ("id", "Mã", 9),
    ("lop", "Lớp", 7),
    ("ho_ten", "Họ tên", 28),
    ("ngay_sinh", "Ngày sinh", 12),
    ("ban_tru", "Bán trú (x)", 11),
    ("ten_ghi_so", "Tên ghi sổ (sửa tay)", 20),
    ("lop_bt", "Lớp tính BT", 11),
    ("an_ghep", "Ăn ghép tại", 11),
    ("cat_thu6", "Cắt thứ 6 (x)", 12),
    ("vang_tu", "Vắng dài hạn từ", 15),
    ("vang_den", "Vắng dài hạn đến", 16),
    ("bt_tu", "Ăn BT từ ngày", 14),
    ("cat_tu", "Cắt BT từ ngày", 14),
    ("ghi_chu", "Ghi chú", 30),
]
BOOL_FIELDS = {"ban_tru", "cat_thu6"}
DATE_FIELDS = {"vang_tu", "vang_den", "bt_tu", "cat_tu"}
LOP_COLS = [("code", "Lớp", 8), ("khoi", "Khối", 7), ("ma_ngan", "Mã ngắn", 10), ("gvcn", "GVCN", 28)]
LOG_COLS = [("Ngày", 12), ("Lớp", 7), ("Đã báo", 8), ("Sĩ số", 7), ("Vắng", 7), ("Tên vắng", 60)]

DEFAULT_SETTINGS = {
    "so_folder": "",
    "so_file": "",
    "cdhs_file": "",
    "vung_in_1": "",  # để trống = tự tính (A1 tới hết phần chữ ký)
    "vung_in_2": "",
}


def _truthy(v) -> bool:
    return fold(str(v or "")) in ("x", "co", "1", "true", "yes", "v", "ban tru")


@dataclass
class AppData:
    path: Path
    roster: Roster = field(default_factory=Roster)
    settings: dict[str, str] = field(default_factory=lambda: dict(DEFAULT_SETTINGS))
    log: dict[tuple[str, str], list] = field(default_factory=dict)  # (ngày, lớp) -> hàng nhật ký

    # ---------- đọc ----------
    @classmethod
    def load(cls, path: str | Path) -> AppData:
        path = Path(path)
        data = cls(path)
        if not path.exists():
            return data
        wb = openpyxl.load_workbook(str(path), data_only=True)
        if "HocSinh" in wb.sheetnames:
            ws = wb["HocSinh"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not any(row):
                    continue
                vals = dict(zip([c[0] for c in HS_COLS], list(row) + [None] * len(HS_COLS), strict=False))
                if not vals.get("ho_ten"):
                    continue
                s = Student(
                    id=clean_spaces(str(vals["id"] or "")) or data.roster.next_id(),
                    lop=norm_class(str(vals["lop"] or "")),
                    ho_ten=clean_spaces(str(vals["ho_ten"])),
                    ngay_sinh=fmt_date(parse_date(vals["ngay_sinh"])) or clean_spaces(str(vals["ngay_sinh"] or "")),
                )
                for f in BOOL_FIELDS:
                    setattr(s, f, _truthy(vals.get(f)))
                for f in DATE_FIELDS:
                    setattr(s, f, parse_date(vals.get(f)))
                s.ten_ghi_so = clean_spaces(str(vals.get("ten_ghi_so") or ""))
                s.lop_bt = norm_class(str(vals.get("lop_bt") or ""))
                s.an_ghep = clean_spaces(str(vals.get("an_ghep") or "")).replace(",", ".")
                s.ghi_chu = clean_spaces(str(vals.get("ghi_chu") or ""))
                data.roster.students.append(s)
        if "Lop" in wb.sheetnames:
            for row in wb["Lop"].iter_rows(min_row=2, values_only=True):
                code = norm_class(str(row[0] or "")) if row else ""
                if code:
                    data.roster.classes[code] = ClassInfo(
                        code, gvcn=clean_spaces(str(row[3] or "")), ma_ngan=clean_spaces(str(row[2] or ""))
                    )
        if "CaiDat" in wb.sheetnames:
            for row in wb["CaiDat"].iter_rows(min_row=2, values_only=True):
                if row and row[0]:
                    data.settings[str(row[0])] = "" if row[1] is None else str(row[1])
        if "NhatKy" in wb.sheetnames:
            for row in wb["NhatKy"].iter_rows(min_row=2, values_only=True):
                if row and row[0] and row[1]:
                    d = parse_date(row[0])
                    if d:
                        data.log[(fmt_date(d), str(row[1]))] = list(row)
        return data

    # ---------- ghi ----------
    def save(self, backup_dir: Path | None = None) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "HocSinh"
        _header(ws, [(t, w) for _, t, w in HS_COLS])
        for s in self.roster.students:
            row = []
            for f, _, _ in HS_COLS:
                v = getattr(s, f)
                if f in BOOL_FIELDS:
                    v = "x" if v else ""
                elif f in DATE_FIELDS:
                    v = fmt_date(v)
                row.append(v)
            ws.append(row)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HS_COLS))}{max(1, ws.max_row)}"

        wl = wb.create_sheet("Lop")
        _header(wl, [(t, w) for _, t, w in LOP_COLS])
        for code in sorted(self.roster.classes, key=class_sort_key):
            c = self.roster.classes[code]
            wl.append([c.code, c.khoi, c.ma_ngan, c.gvcn])

        wk = wb.create_sheet("NhatKy")
        _header(wk, LOG_COLS)
        for key in sorted(self.log, key=lambda k: (parse_date(k[0]) or dt.date.min, class_sort_key(k[1]))):
            wk.append(self.log[key])

        wc = wb.create_sheet("CaiDat")
        _header(wc, [("Khóa", 14), ("Giá trị", 80)])
        for k, v in self.settings.items():
            wc.append([k, v])

        safe_save(wb, self.path, backup_dir)

    def set_log(self, day: dt.date, code: str, reported: bool, si_so: int, names: list[str]) -> None:
        self.log[(fmt_date(day), code)] = [fmt_date(day), code, "x" if reported else "", si_so, len(names), ", ".join(names)]

    def reported_flags(self, day: dt.date) -> dict[str, bool]:
        key = fmt_date(day)
        return {c: _truthy(row[2]) for (d, c), row in self.log.items() if d == key}


def _header(ws, cols: list[tuple[str, int]]) -> None:
    ws.append([t for t, _ in cols])
    fill = PatternFill("solid", fgColor="FFDDEBF7")
    for i, (_, w) in enumerate(cols, 1):
        cell = ws.cell(1, i)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def short_codes(codes: list[str]) -> dict[str, str]:
    """Mã ngắn theo thứ tự trong khối: 1D15 -> 1.1, 1D16 -> 1.2…"""
    out: dict[str, str] = {}
    by_khoi: dict[str, list[str]] = {}
    for c in sorted(codes, key=class_sort_key):
        by_khoi.setdefault(c[0], []).append(c)
    for k, lst in by_khoi.items():
        for i, c in enumerate(lst, 1):
            out[c] = f"{k}.{i}"
    return out


def default_data_path() -> Path:
    return Path.home() / "Documents" / "DiemDanhBanTru" / "DuLieu_BanTru.xlsx"
