import datetime as dt

import openpyxl
import pytest
from openpyxl.styles import Font

from app.core.day import DayState
from app.excel import sobook
from app.excel.database import AppData, short_codes
from app.excel.files import FileLockedError, lock_file, safe_save
from app.excel.roster_import import apply_changes, diff_roster, flag_mismatches, read_cdhs

DAY28 = dt.date(2026, 10, 28)


def make_book(path):
    """Sổ giả cùng cấu trúc sổ thật: tiêu đề, khối, tổng, chữ ký, bảng suất ăn có công thức."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "28"
    tnr = Font(name="Times New Roman", size=11)
    ws["A1"] = "UBND PHƯỜNG X"
    ws["A2"] = "TRƯỜNG TIỂU HỌC X"
    ws["A3"] = "ĐIỂM DANH BÁN TRÚ NGÀY 28/10/2026"
    ws["A3"].font = Font(name="Times New Roman", size=14, bold=True)
    ws["A4"] = "Người lấy sĩ số: Giáo Vụ"
    for c, h in zip("ABCDEFG", ["STT", "LỚP", "GIÁO VIÊN CHỦ NHIỆM", "SỈ số", "Vắng", "HIỆN DIỆN", "TÊN HỌC SINH VẮNG"], strict=True):
        ws[f"{c}5"] = h
    rows = [(6, "1D15", 6, 1, "Tuấn Anh"), (7, "1D16", 5, 0, None), (8, "1D19", 4, 0, None), (9, "KHỐI 1",),
            (10, "2D17", 4, 0, None), (11, "KHỐI 2",), (12, "4D19", 2, 0, None), (13, "4D21", 3, 1, "Vũ Khoa"),
            (14, "KHỐI 4",)]
    for item in rows:
        r = item[0]
        if len(item) == 2:
            ws[f"A{r}"] = item[1]
            continue
        _, code, si, vang, names = item
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"] = r - 5, code, "GV " + code, si, vang
        ws[f"F{r}"] = f"=D{r}-E{r}"
        ws[f"G{r}"] = names
        ws[f"B{r}"].font = tnr
    ws["H13"] = "Vũ Khoa vắng tới khi có thông báo mới"
    ws["A15"] = "TỔNG CỘNG"
    ws["G16"] = "Phường X, ngày     tháng 10 năm 2026  "
    ws["I20"] = "Suất ăn"
    ws["I21"], ws["J21"] = "=F10", "hiến"
    ws["G21"] = "=SUM(I21:I25)"
    ws["M22"] = "vắng"
    ws["I23"], ws["J23"], ws["L23"], ws["O23"] = "=F6+L23", "1.1 ghép", "=O23-M23", 2
    ws["I24"], ws["J24"], ws["L24"], ws["O24"] = "=F7+L24", "1.2 ghép", "=O24-M24", 2
    ws["I25"], ws["J25"], ws["L25"], ws["O25"], ws["M25"] = "=F12+L25", "4.2 ghép 4.4", "=O25-M25", 3, 1
    ws.page_setup.orientation = "landscape"
    wb.save(path)
    return path


@pytest.fixture
def book(tmp_path):
    return make_book(tmp_path / "SO DIEM DANH T10.xlsx")


def test_analyze_layout(book):
    wb = sobook.load_book(book)
    info = sobook.analyze(wb["28"])
    assert list(info.class_rows) == ["1D15", "1D16", "1D19", "2D17", "4D19", "4D21"]
    assert info.total_row == 15 and info.title_cell == "A3" and info.sign_cell == "G16"
    m = info.meal
    assert m.ghep_classes == ["1D19", "4D21"]
    assert [(r.label, r.classes, r.code) for r in m.rows] == [
        ("hiến", ["2D17"], None),
        ("1.1 ghép", ["1D15"], "1.1"),
        ("1.2 ghép", ["1D16"], "1.2"),
        ("4.2 ghép 4.4", ["4D19"], "4.2"),
    ]
    assert m.o_values == {"1.1": 2, "1.2": 2, "4.2": 3}
    assert m.area == "G20:O25"
    assert sobook.book_month(wb) == (10, 2026)


def test_create_write_and_read_back(book, roster):
    wb = sobook.load_book(book)
    tpl = sobook.template_sheet(wb, dt.date(2026, 10, 29))
    ws = sobook.create_day_sheet(wb, dt.date(2026, 10, 29), tpl)
    assert wb.sheetnames == ["28", "29"]
    info = sobook.analyze(ws)
    assert ws["A3"].value == "ĐIỂM DANH BÁN TRÚ NGÀY 29/10/2026"
    assert ws["G6"].value is None and ws["E6"].value == 0 and ws["M25"].value is None
    assert ws["H13"].value.startswith("Vũ Khoa")  # ghi chú giữ nguyên
    assert ws["B6"].font.name == "Times New Roman"  # định dạng giữ nguyên
    assert ws["I23"].value == "=F6+L23"  # công thức giữ nguyên

    for s in roster.students:
        if s.lop == "1D19":
            s.an_ghep = "1.1" if s.ho_ten in ("Châu Bảo Ngọc", "Phùng Gia Hân") else "1.2"
        if s.lop == "4D21":
            s.an_ghep = "4.2"
    st = DayState.new(dt.date(2026, 10, 29), roster, list(info.class_rows))
    st.set_absent("1D15", roster.students[3].id, True)  # Tuấn Anh
    st.set_absent("1D19", roster.students[15].id, True)  # Tạ Đức Anh -> 1.2
    st.cd("1D15").note = "ghi chú thử"
    summary = st.meals(info.meal)
    assert summary.balanced, summary.warnings
    sobook.write_day(info, st, summary)
    sobook.select_only(wb, ws)
    sobook.save_book(wb, book)

    wb2 = sobook.load_book(book)
    ws2 = wb2["29"]
    rec = sobook.read_records(sobook.analyze(ws2))
    assert (rec["1D15"].si_so, rec["1D15"].vang, rec["1D15"].names, rec["1D15"].note) == (6, 1, "Tuấn Anh", "ghi chú thử")
    assert rec["1D19"].names == "Đức Anh"
    assert ws2["M24"].value == 1 and ws2["M23"].value is None and ws2["O24"].value == 2
    assert [s.sheet_view.tabSelected for s in wb2.worksheets] == [False, True]
    assert (book.parent / "SaoLuu").is_dir()  # đã sao lưu bản cũ


def test_new_month_book(book, tmp_path):
    new = sobook.month_file_name(book, 11)
    assert new.name == "SO DIEM DANH T11.xlsx"
    sobook.new_month_book(book, dt.date(2026, 11, 2), new)
    wb = sobook.load_book(new)
    assert wb.sheetnames == ["2"]
    ws = wb["2"]
    assert ws["A3"].value == "ĐIỂM DANH BÁN TRÚ NGÀY 02/11/2026"
    assert "tháng 11 năm 2026" in ws["G16"].value
    assert ws["G6"].value is None and ws["M25"].value is None
    assert sobook.find_month_file(tmp_path, 11) == new


def test_safe_save_refuses_when_open_in_excel(book):
    lock_file(book).write_text("x")
    wb = sobook.load_book(book)
    with pytest.raises(FileLockedError):
        safe_save(wb, book)


def test_app_data_roundtrip(tmp_path, roster):
    roster.students[0].cat_thu6 = True
    roster.students[0].vang_tu = dt.date(2026, 10, 1)
    roster.students[1].an_ghep = "1.2"
    data = AppData(tmp_path / "DuLieu_BanTru.xlsx", roster)
    data.settings["so_folder"] = str(tmp_path)
    data.set_log(DAY28, "1D15", True, 6, ["Tuấn Anh"])
    data.save()
    back = AppData.load(tmp_path / "DuLieu_BanTru.xlsx")
    s0, s1 = back.roster.students[0], back.roster.students[1]
    assert (s0.ho_ten, s0.cat_thu6, s0.vang_tu, s0.ban_tru) == ("Kha Bảo Ngọc", True, dt.date(2026, 10, 1), True)
    assert s1.an_ghep == "1.2"
    assert len(back.roster.students) == len(roster.students)
    assert back.settings["so_folder"] == str(tmp_path)
    assert back.reported_flags(DAY28) == {"1D15": True}
    assert short_codes(["1D16", "1D15", "2D17"]) == {"1D15": "1.1", "1D16": "1.2", "2D17": "2.1"}


def _write_fake_cdhs_xls(path):
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("CDAN")
    red = xlwt.easyxf("font: colour red")
    for c, h in enumerate(["TT", "Tên lớp học", "Họ tên", "Ngày sinh", "Phân loại HS"]):
        ws.write(0, c, h)
    data = [("Lớp 1D15", "Kha Bảo Ngọc", "02/03/2020", True, "Bán trú"),
            ("Lớp 1D15", "Mạc Trần Quỳnh Anh", "01/01/2020", False, "Không"),
            ("Lớp 1D16", "Lý Gia Hưng", "13/03/2020", True, "Không")]
    for r, (lop, name, dob, is_red, pl) in enumerate(data, 1):
        ws.write(r, 0, r)
        ws.write(r, 1, lop)
        ws.write(r, 2, name, red if is_red else xlwt.easyxf(""))
        ws.write(r, 3, dob)
        ws.write(r, 4, pl)
    wb.save(str(path))


def test_read_cdhs_xls_red_font(tmp_path):
    p = tmp_path / "cdhs.xls"
    _write_fake_cdhs_xls(p)
    rows = read_cdhs(p)
    assert [(r.lop, r.ho_ten, r.red) for r in rows] == [
        ("1D15", "Kha Bảo Ngọc", True), ("1D15", "Mạc Trần Quỳnh Anh", False), ("1D16", "Lý Gia Hưng", True)]
    assert [r.ho_ten for r in flag_mismatches(rows)] == ["Lý Gia Hưng"]


def test_read_cdhs_xlsx_and_diff(tmp_path, roster):
    p = tmp_path / "cdhs.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["TT", "Tên lớp học", "Họ tên", "Ngày sinh"])
    ws.append([1, "Lớp 1D15", "Kha Bảo Ngọc", "02/03/2020"])  # thôi ăn BT (không đỏ)
    ws.append([2, "Lớp 1D16", "Mạc Trần Quỳnh Anh", "01/01/2020"])  # đổi lớp
    ws.append([3, "Lớp 1D16", "Em Mới", "05/05/2020"])
    ws["C4"].font = Font(color="FFFF0000")
    wb.save(p)
    rows = read_cdhs(p)
    assert [r.red for r in rows] == [False, False, True]
    changes = diff_roster(roster, rows)
    kinds = sorted(c.kind for c in changes)
    assert kinds.count("removed") == len(roster.students) - 2
    assert {"new", "class", "boarding"} <= set(kinds)
    n_before = len(roster.students)
    apply_changes(roster, changes)  # mặc định không xóa HS vắng mặt trong file mới
    assert len(roster.students) == n_before + 1
    assert roster.students[0].ban_tru is False and roster.students[5].lop == "1D16"
